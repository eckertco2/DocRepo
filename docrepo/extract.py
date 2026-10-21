"""Pulling the plain text out of a file, for search and for the differences in document versions.

Files are only ever opened to read. The text is cut at ``extracted_text_cap_chars``
(settings.json). A file DocRepo can't read is still recorded; its status says why.
This isn't Extract, the later processing into the knowledge base.
"""
from __future__ import annotations

import hashlib
import logging
import os
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path

log = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".log", ".json", ".xml", ".ini"}
#: What each kind of file is read with, for the help text and SPEC.
READERS = {".pdf": "pypdf", ".docx": "python-docx", ".pptx": "zipfile", ".xlsx": "openpyxl", ".xlsm": "openpyxl", ".eml": "email"}


@dataclass(frozen=True)
class Extracted:
    """A file's text and how it went: ok, partial (cut at the limit), none (no text found or not a kind DocRepo reads) or error."""

    text: str
    status: str
    note: str = ""


@dataclass(frozen=True)
class FileInfo:
    size: int
    modified: str
    sha256: str


def file_info(path: Path) -> FileInfo:
    """A file's size, modified time and SHA-256, for noticing later that it changed."""
    stat = path.stat()
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            digest.update(block)
    return FileInfo(stat.st_size, datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"), digest.hexdigest())


def decode(data: bytes) -> str:
    """Text from bytes: UTF-8 (with or without a BOM), UTF-16 with a BOM, or else Windows-1252."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()


# ---- .pptx, read with the standard library ----
# A .pptx is a zip of XML parts. Reading it here, rather than with python-pptx, keeps out python-pptx's
# dependencies, Pillow among them, whose bundled licenses a company scan flagged.
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
#: The largest XML part read from a .pptx, so a file built to unpack to gigabytes is refused.
MAX_PART_BYTES = 64 * 1024 * 1024


def _has(z: zipfile.ZipFile, name: str) -> bool:
    try:
        z.getinfo(name)
    except KeyError:
        return False
    return True


def _xml(z: zipfile.ZipFile, name: str) -> ET.Element:
    """An XML part of the zip. A part with a DTD is refused: Office never writes one, and it's how entity tricks work."""
    if z.getinfo(name).file_size > MAX_PART_BYTES:
        raise ValueError(f"{name} is too big to read")
    data = z.read(name)
    if b"<!DOCTYPE" in data:
        raise ValueError(f"{name} has a DTD")
    return ET.fromstring(data)


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """A part's relationships: their IDs, each with its type and the name of the part it points to."""
    folder, name = posixpath.split(part)
    rels_name = posixpath.join(folder, "_rels", name + ".rels")
    if not _has(z, rels_name):
        return {}
    out = {}
    for rel in _xml(z, rels_name).iter(f"{_REL}Relationship"):
        if rel.get("TargetMode") == "External":
            continue
        target = rel.get("Target", "")
        full = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id", "")] = (rel.get("Type", ""), full)
    return out


def _text_body(body: ET.Element | None) -> str:
    """A text box's text: a line for each paragraph, and a line break inside one kept as a new line."""
    if body is None:
        return ""
    lines = []
    for paragraph in body.iter(f"{_A}p"):
        bits = [(el.text or "") if el.tag == f"{_A}t" else "\n" for el in paragraph.iter() if el.tag in (f"{_A}t", f"{_A}br")]
        lines.append("".join(bits))
    return "\n".join(lines)


def _shapes(tree: ET.Element, parts: list[str]) -> None:
    """The text of a slide's shapes, in order: text boxes, tables row by row, and the shapes inside groups."""
    for shape in tree:
        if shape.tag == f"{_P}sp":
            parts.append(_text_body(shape.find(f"{_P}txBody")))
        elif shape.tag == f"{_P}grpSp":
            _shapes(shape, parts)
        elif shape.tag == f"{_P}graphicFrame":
            for row in shape.iter(f"{_A}tr"):
                parts.append(" | ".join(_text_body(cell.find(f"{_A}txBody")).strip() for cell in row.iter(f"{_A}tc")))


def _notes(z: zipfile.ZipFile, part: str) -> str:
    """A notes page's speaker notes: the text of its body placeholder."""
    for shape in _xml(z, part).iter(f"{_P}sp"):
        placeholder = shape.find(f"{_P}nvSpPr/{_P}nvPr/{_P}ph")
        if placeholder is not None and placeholder.get("type") == "body":
            return _text_body(shape.find(f"{_P}txBody")).strip()
    return ""


def _pptx(path: Path) -> str:
    """Each slide's text in the deck's order, under "Slide 1" and so on, then its speaker notes as "Notes: …"."""
    parts: list[str] = []
    with zipfile.ZipFile(path) as z:
        main = next((t for kind, t in _rels(z, "").values() if kind.endswith("/officeDocument")), "ppt/presentation.xml")
        slides = _rels(z, main)
        for number, slide_id in enumerate(_xml(z, main).iter(f"{_P}sldId"), 1):
            part = slides.get(slide_id.get(f"{_R}id", ""), ("", ""))[1]
            if not _has(z, part):
                continue
            parts.append(f"Slide {number}")
            tree = _xml(z, part).find(f"{_P}cSld/{_P}spTree")
            if tree is not None:
                _shapes(tree, parts)
            notes = next((t for kind, t in _rels(z, part).values() if kind.endswith("/notesSlide") and _has(z, t)), None)
            if notes and (text := _notes(z, notes)):
                parts.append(f"Notes: {text}")
    return "\n".join(p for p in parts if p).strip()


def _xlsx(path: Path, cap: int) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    size = 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                if any(values):
                    line = " | ".join(values).rstrip(" |")
                    parts.append(line)
                    size += len(line) + 1
                if size > cap:  # big sheets stop early; the text is cut at the cap anyway
                    return "\n".join(parts)
    finally:
        workbook.close()
    return "\n".join(parts).strip()


def _eml(path: Path) -> str:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    head = [f"{k}: {message[k]}" for k in ("From", "To", "Date", "Subject") if message[k]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    return "\n".join(head) + "\n\n" + text.strip()


def extract_text(path: Path, cap: int) -> Extracted:
    """The plain text of a file, read-only. Never raises: a failure comes back with status "error"."""
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES:
            text = decode(path.read_bytes())
        elif suffix == ".pdf":
            text = _pdf(path)
        elif suffix == ".docx":
            text = _docx(path)
        elif suffix == ".pptx":
            text = _pptx(path)
        elif suffix in (".xlsx", ".xlsm"):
            text = _xlsx(path, cap)
        elif suffix == ".eml":
            text = _eml(path)
        else:
            return Extracted("", "none", f"DocRepo doesn't read {suffix or 'extensionless'} files")
    except Exception as e:  # noqa: BLE001 - a bad file must never stop a save
        log.warning("Couldn't read the text of %s: %s", path.name, type(e).__name__)
        return Extracted("", "error", f"couldn't read it ({type(e).__name__})")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        note = "no text found; it may be a scanned document" if suffix == ".pdf" else "no text found"
        return Extracted("", "none", note)
    if len(text) > cap:
        return Extracted(text[:cap], "partial", f"cut at {cap} characters")
    return Extracted(text, "ok")


def to_unc(path: str) -> str:
    """A path on a mapped network drive (H:\\x) as its network path (\\\\server\\share\\x), so it survives drive-letter changes.

    Asks Windows which share the letter is mapped to; nothing is sent over the network.
    """
    drive = os.path.splitdrive(path)[0]
    if os.name != "nt" or len(drive) != 2 or drive[1] != ":":
        return path
    import ctypes
    from ctypes import wintypes

    if ctypes.windll.kernel32.GetDriveTypeW(f"{drive}\\") != 4:  # DRIVE_REMOTE
        return path
    buffer = ctypes.create_unicode_buffer(1024)
    length = wintypes.DWORD(len(buffer))
    if ctypes.windll.mpr.WNetGetConnectionW(drive, buffer, ctypes.byref(length)) != 0:
        return path
    return buffer.value.rstrip("\\") + path[2:]
