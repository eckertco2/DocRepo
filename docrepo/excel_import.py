"""The one-time import of the project list from Excel: each row becomes a top-level task.

The workbook is opened read-only (openpyxl's read-only mode, with cached values,
so formulas and macros never run) and is never changed. The browse tool's form
shows its columns with a proposed mapping to task fields, which the owner
changes as needed, and a preview; nothing is written until Import.

- A column can fill a task field, add "Header: value" lines to the description
  (so nothing is lost), stand for one label (a yes or no column), or be left out.
- A State column's values each map to a state; without one, every task starts
  in the state the owner picks (In progress by default).
- Names in Lead and Assignees columns are matched to the people list by name or
  alias; the rest are added as new people. Label names that don't exist yet
  become new labels.
- Dates are read from date cells, or from text the way Parse and populate reads
  date phrases ("Q1 2027", "10/30/2026"). Text that isn't a date goes into the
  description, and the preview says so.
- A row whose title is already an imported task's title is skipped, so running
  the import twice doesn't make doubles.

Each task records "Excel import" as what created it, so its story begins
"Imported from the Excel tracker". The log records counts, never the workbook's contents.
"""
from __future__ import annotations

import json
import logging
import re
import sqlite3
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from docrepo import db, people, tasks
from docrepo.config import Settings
from docrepo.dates import stamp
from docrepo.db import STATES, transaction
from docrepo.model import Model
from docrepo.pages import write_pages
from docrepo.parse import find_dates
from docrepo.safe_fs import SafeFS

log = logging.getLogger(__name__)

CREATED_BY = "Excel import"
STATE_KEY = "excel_import"
MAX_ROWS = 5000
SAMPLES = 3

#: What a column can fill, in the order the form lists them.
FIELDS = (
    ("title", "Title"), ("description", "Description"), ("extra", "Add to the description"), ("current_status", "Current status"),
    ("state", "State"), ("start", "Start"), ("due", "Due"), ("deadline", "Deadline"), ("deadline_desc", "Deadline description"),
    ("leads", "Lead"), ("assignees", "Assignees"), ("labels", "Labels"), ("label", "A label (yes or no)"), ("urgent", "Urgent"),
    ("lead_hours", "Lead hours"), ("assignee_hours", "Assignee hours"), ("ignore", "Leave out"),
)
FIELD_NAMES = dict(FIELDS)
#: Fields only one column can fill.
SINGLE = ("title", "description", "current_status", "state", "start", "due", "deadline", "deadline_desc", "urgent", "lead_hours",
          "assignee_hours")
DATE_FIELDS = ("start", "due", "deadline")

#: Header words that suggest a field, tried in order. A header matches when one of these is a whole word in it.
_GUESSES = (
    ("deadline_desc", ("deadline description", "deadline event", "event")),
    ("deadline", ("deadline", "hard date", "must be done")),
    ("due", ("due", "due date", "target", "target date", "finish", "completion", "end date", "eta")),
    ("start", ("start", "start date", "begin", "kick off", "kickoff")),
    ("lead_hours", ("lead hours",)),
    ("assignee_hours", ("assignee hours", "hours", "effort", "estimate")),
    ("current_status", ("status update", "current status", "update", "latest", "comments", "comment", "notes", "progress", "status")),
    ("state", ("state", "stage", "phase")),
    ("description", ("description", "scope", "details", "summary", "objective", "purpose")),
    ("leads", ("lead", "owner", "responsible", "manager", "pm", "champion", "sponsor")),
    ("assignees", ("assignee", "assignees", "assigned", "team", "resources", "resource", "support", "members", "who")),
    ("labels", ("label", "labels", "category", "type", "program", "programme", "area", "unit", "line")),
    ("urgent", ("urgent", "priority", "critical")),
    # Last, so "Project lead" or "Owner name" isn't taken for the title.
    ("title", ("project", "project name", "name", "title", "task", "item")),
)
_TRUE = {"yes", "y", "x", "true", "1", "✓", "✔", "v"}
_FALSE = {"", "no", "n", "false", "0", "-", "–"}
_URGENT = _TRUE | {"high", "urgent", "critical", "top", "p1", "1 - high", "h"}
_SPLIT_PEOPLE = re.compile(r"\s*(?:[,;/&\n+]|\band\b)\s*", re.IGNORECASE)
_SPLIT_LABELS = re.compile(r"\s*[,;/\n]\s*")


class WorkbookError(Exception):
    """The workbook can't be read, or the mapping can't be used."""


# ---- reading the workbook ----
def sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    book = _open(path, load_workbook)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def _open(path: Path, load_workbook):
    try:
        return load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    except PermissionError as e:
        raise WorkbookError(f"{path.name} can't be opened. If it's open in Excel, close it, or save a copy and import that.") from e
    except (OSError, ValueError, KeyError) as e:
        raise WorkbookError(f"{path.name} can't be read as an Excel workbook ({type(e).__name__}).") from e
    except Exception as e:  # noqa: BLE001 - openpyxl raises many kinds of error for a damaged file
        raise WorkbookError(f"{path.name} can't be read as an Excel workbook ({type(e).__name__}).") from e


def read_sheet(path: Path, sheet: str) -> list[list[object]]:
    """The sheet's rows as lists of cell values, without trailing empty cells, at most MAX_ROWS."""
    from openpyxl import load_workbook

    book = _open(path, load_workbook)
    try:
        if sheet not in book.sheetnames:
            raise WorkbookError(f"There's no sheet called {sheet} in {path.name}.")
        rows: list[list[object]] = []
        for values in book[sheet].iter_rows(values_only=True):
            row = list(values)
            while row and _blank(row[-1]):
                row.pop()
            rows.append(row)
            if len(rows) >= MAX_ROWS:
                break
        while rows and not rows[-1]:
            rows.pop()
        return rows
    finally:
        book.close()


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def text_of(value: object) -> str:
    """A cell as text: dates as YYYY-MM-DD, whole numbers without .0."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def guess_header_row(rows: list[list[object]]) -> int:
    """The row of headers (0 is the first): the first with two or more text cells, followed by a row with something in it."""
    for i, row in enumerate(rows[:30]):
        texts = [v for v in row if isinstance(v, str) and v.strip()]
        if len(texts) >= 2 and any(not _blank(v) for r in rows[i + 1:i + 3] for v in r):
            return i
    return 0


# ---- the mapping ----
@dataclass
class Column:
    index: int
    header: str
    samples: list[str]
    field: str = "extra"
    label: str = ""  # for "label": the label a yes in this column puts on the task


def _words(header: str) -> str:
    return " " + re.sub(r"[^a-z0-9]+", " ", header.lower()).strip() + " "


def guess_field(header: str, values: list[object]) -> str:
    """The field a column most likely fills, from its header, or from its values for a yes or no column."""
    texts = [text_of(v).lower() for v in values if not _blank(v)]
    if not texts:
        return "ignore"
    words = _words(header)
    by_header = next((key for key, hints in _GUESSES if any(f" {h} " in words for h in hints)), None)
    yes_no = all(t in _TRUE | _FALSE for t in texts)
    if yes_no and by_header in (None, "labels", "title"):  # "Line 1" holding x or yes is one label
        return "label"
    return by_header or "extra"


def propose(rows: list[list[object]], header_row: int) -> list[Column]:
    """Each column with its header, a few sample values and the field it most likely fills. Each single field goes to one column."""
    headers = rows[header_row] if header_row < len(rows) else []
    body = rows[header_row + 1:]
    width = max([len(headers), *[len(r) for r in body]] or [0])
    out: list[Column] = []
    taken: set[str] = set()
    for i in range(width):
        header = text_of(headers[i]) if i < len(headers) else ""
        values = [r[i] for r in body if i < len(r)]
        guess = guess_field(header, values)
        if guess in SINGLE and guess in taken:
            guess = "extra"
        taken.add(guess)
        samples = [text_of(v) for v in values if not _blank(v)][:SAMPLES]
        out.append(Column(i, header or f"Column {i + 1}", samples, guess, header if guess == "label" else ""))
    if "title" not in taken:  # with no header that names it, the first text column is the title
        first = next((c for c in out if c.field in ("extra", "description") and c.samples), None)
        if first:
            first.field = "title"
    return out


def guess_state(value: str) -> str:
    v = value.lower()
    if re.search(r"\b(done|complete|completed|closed|finished|cancel+ed|dropped|delivered)\b", v):
        return "Closed"
    if re.search(r"\b(hold|parked|paused|waiting|deferred|on ice|backlog)\b", v):
        return "Parked"
    if re.search(r"\b(not started|new|planned|proposed|idea|future|to do|todo)\b", v):
        return "Not started"
    return "In progress"


def state_values(rows: list[list[object]], header_row: int, column: Column | None) -> list[tuple[str, int, str]]:
    """The State column's values: (value, how many rows, the proposed state), most common first."""
    if column is None:
        return []
    counts = Counter(text_of(r[column.index]) for r in rows[header_row + 1:] if column.index < len(r) and not _blank(r[column.index]))
    return [(v, n, guess_state(v)) for v, n in counts.most_common()]


@dataclass
class Plan:
    columns: list[Column]
    header_row: int
    default_state: str = "In progress"
    state_map: dict[str, str] = field(default_factory=dict)  # a State column's value -> state

    def column_for(self, key: str) -> Column | None:
        return next((c for c in self.columns if c.field == key), None)

    def check(self) -> None:
        if self.column_for("title") is None:
            raise WorkbookError("Choose the column that holds each project's name (Title).")
        for key in SINGLE:
            if sum(1 for c in self.columns if c.field == key) > 1:
                raise WorkbookError(f"Only one column can fill {FIELD_NAMES[key]}.")
        if self.default_state not in STATES:
            raise WorkbookError(f"Unknown state {self.default_state}.")
        for c in self.columns:
            if c.field == "label" and not c.label.strip():
                raise WorkbookError(f"Name the label that column {c.header} stands for.")


# ---- the preview ----
@dataclass
class Row:
    line: int  # the row's number in Excel
    title: str
    fields: dict[str, object]
    leads: list[str]  # names, as in the workbook
    assignees: list[str]
    labels: list[str]
    problems: list[str] = field(default_factory=list)
    skip: str = ""  # why it isn't imported, if it isn't


@dataclass
class Preview:
    rows: list[Row]
    new_people: list[str]
    new_labels: list[str]

    @property
    def to_import(self) -> list[Row]:
        return [r for r in self.rows if not r.skip]


def _date(value: object, today: date) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # a date stored as Excel's day number
        return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
    text = text_of(value)
    if re.fullmatch(r"\d{4}-\d\d-\d\d.*", text):
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError:
            pass
    hits = find_dates(text, today) if text else []
    return hits[0].date.isoformat() if hits else None


def _hours(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hours?)?\s*", text_of(value), re.IGNORECASE)
    return float(m.group(1)) if m else None


def _truthy(value: object, words: set[str]) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return text_of(value).lower() in words


class _Names:
    """Matches names to the people list by name or alias, any case; counts the ones that will be new.
    A retired person is recognized but left off, as Parse and populate does."""

    def __init__(self, model: Model) -> None:
        self.by_name: dict[str, str] = {}
        self.retired: dict[str, str] = {}
        for p in model.people.values():
            for n in (p.name, *p.aliases):
                (self.by_name if p.status == "active" else self.retired).setdefault(n.lower(), p.name)
        self.new: list[str] = []

    def resolve(self, cell: object, problems: list[str]) -> list[str]:
        out = []
        for name in _SPLIT_PEOPLE.split(text_of(cell)):
            name = name.strip(" .")
            if not name:
                continue
            if name.lower() not in self.by_name and name.lower() in self.retired:
                problems.append(f"{name} isn't added: {self.retired[name.lower()]} is retired on the people list.")
                continue
            if name.lower() not in self.by_name and not any(n.lower() == name.lower() for n in self.new):
                self.new.append(name)
            out.append(name)
        return list(dict.fromkeys(out))


def preview(rows: list[list[object]], plan: Plan, model: Model, today: date) -> Preview:
    """What Import would create, row by row, with what it can't read. Nothing is written."""
    plan.check()
    names = _Names(model)
    known_labels = {lb.name.lower(): lb.name for lb in model.labels}
    new_labels: list[str] = []
    done_titles = {t.title.strip().lower() for t in model.tasks.values() if t.created_by == CREATED_BY}
    seen: dict[str, int] = {}
    out: list[Row] = []
    for offset, cells in enumerate(rows[plan.header_row + 1:]):
        line = plan.header_row + 2 + offset

        def cell(c: Column) -> object:
            return cells[c.index] if c.index < len(cells) else None

        title_col = plan.column_for("title")
        title = " ".join(text_of(cell(title_col)).split())
        if not title:
            continue  # an empty row, or a row of notes under the list
        row = Row(line, title[:200], {}, [], [], [])
        extras: list[str] = []
        for c in plan.columns:
            value = cell(c)
            if c.field in ("ignore", "title") or (_blank(value) and c.field != "label"):
                continue
            text = text_of(value)
            if c.field == "description":
                row.fields["description"] = text
            elif c.field == "extra":
                extras.append(f"{c.header}: {text}")
            elif c.field == "current_status":
                row.fields["current_status"] = text
            elif c.field == "deadline_desc":
                row.fields["deadline_desc"] = text
            elif c.field == "state":
                row.fields["state"] = plan.state_map.get(text, guess_state(text))
            elif c.field in DATE_FIELDS:
                day = _date(value, today)
                if day:
                    row.fields[c.field] = day
                else:
                    extras.append(f"{c.header}: {text}")
                    row.problems.append(f"{FIELD_NAMES[c.field]} “{text}” isn't a date, so it's added to the description.")
            elif c.field in ("lead_hours", "assignee_hours"):
                hours = _hours(value)
                if hours is not None and hours >= 0:
                    row.fields[c.field] = hours
                else:
                    extras.append(f"{c.header}: {text}")
                    row.problems.append(f"{FIELD_NAMES[c.field]} “{text}” isn't a number of hours, so it's added to the description.")
            elif c.field in ("leads", "assignees"):
                getattr(row, c.field).extend(n for n in names.resolve(value, row.problems) if n not in getattr(row, c.field))
            elif c.field == "labels":
                for name in _SPLIT_LABELS.split(text):
                    if name and name.lower() not in known_labels and name.lower() not in {n.lower() for n in new_labels}:
                        new_labels.append(name)
                    if name:
                        row.labels.append(known_labels.get(name.lower(), next((n for n in new_labels if n.lower() == name.lower()), name)))
            elif c.field == "label":
                if _truthy(value, _TRUE):
                    name = c.label.strip()
                    if name.lower() not in known_labels and name.lower() not in {n.lower() for n in new_labels}:
                        new_labels.append(name)
                    row.labels.append(known_labels.get(name.lower(), next((n for n in new_labels if n.lower() == name.lower()), name)))
            elif c.field == "urgent":
                row.fields["urgent"] = _truthy(value, _URGENT)
        row.labels = list(dict.fromkeys(row.labels))
        if extras:
            row.fields["description"] = "\n".join(x for x in (row.fields.get("description", ""), "\n".join(extras)) if x).strip()
        row.fields.setdefault("state", plan.default_state)
        if row.fields.get("deadline_desc") and not row.fields.get("deadline"):
            row.problems.append("It has a deadline description but no deadline.")
        key = row.title.lower()
        if key in done_titles:
            row.skip = "Already imported: a task with this title came from the Excel import before."
        elif key in seen:
            row.problems.append(f"Row {seen[key]} has the same title.")
        seen.setdefault(key, line)
        out.append(row)
    return Preview(out, names.new, new_labels)


# ---- importing ----
def run_import(conn: sqlite3.Connection, settings: Settings, data_root: Path, fs: SafeFS, found: Preview, at: datetime,
               source: str = "") -> str:
    """Create the previewed tasks, with their new people and labels, in one transaction. Returns what it did, in a sentence."""
    rows = found.to_import
    if not rows:
        return "Nothing to import: every row is skipped."
    day = at.date().isoformat()
    with transaction(conn):
        model = Model(conn)
        by_name = {}
        for p in model.active_people():
            for n in (p.name, *p.aliases):
                by_name.setdefault(n.lower(), p.id)
        for name in found.new_people:
            if name.lower() not in by_name:
                by_name[name.lower()] = people.add_person(conn, name, added_note=f"Added on {day} by the Excel import. Aliases and hours aren't set yet.")
        existing = {lb.name.lower() for lb in model.labels}
        for name in found.new_labels:
            if name.lower() not in existing:
                people.add_label(conn, name)
                existing.add(name.lower())
        for r in rows:
            f = dict(r.fields)
            tasks.create_task(
                conn, settings, title=r.title, at=at, created_by=CREATED_BY, description=str(f.pop("description", "")),
                leads=[by_name[n.lower()] for n in r.leads], assignees=[by_name[n.lower()] for n in r.assignees], labels=r.labels, **f,
            )
        db.set_state(conn, STATE_KEY, json.dumps({"at": stamp(at), "tasks": len(rows), "file": source}))
    try:
        write_pages(Model(conn), data_root, fs, at.date(), at)
    except OSError as e:
        log.warning("excel import: pages not all written (%s)", type(e).__name__)
    log.info("excel import: %d tasks, %d new people, %d new labels", len(rows), len(found.new_people), len(found.new_labels))
    bits = [f"Imported {len(rows)} task{'s' if len(rows) != 1 else ''} from Excel"]
    if found.new_people:
        bits.append(f"added {len(found.new_people)} {'person' if len(found.new_people) == 1 else 'people'}")
    if found.new_labels:
        bits.append(f"{len(found.new_labels)} label{'s' if len(found.new_labels) != 1 else ''}")
    skipped = len(found.rows) - len(rows)
    return ", ".join(bits) + "." + (f" {skipped} row{'s were' if skipped != 1 else ' was'} skipped." if skipped else "")


def last_import(conn: sqlite3.Connection) -> dict | None:
    raw = db.get_state(conn, STATE_KEY)
    try:
        return json.loads(raw) if raw else None
    except ValueError:
        return None
