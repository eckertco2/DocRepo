"""Pulling text out of files: each kind DocRepo reads, the cap, and files it can't read."""
from __future__ import annotations

import zipfile
from pathlib import Path

from docrepo.extract import extract_text, file_info, to_unc

CAP = 10_000


def make_pdf(path: Path, line: str) -> None:
    """A one-page PDF with a line of text, written by hand (no PDF-writing package is approved)."""
    stream = f"BT /F1 18 Tf 20 100 Td ({line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def test_plain_text_in_any_common_encoding(tmp_path):
    for name, data in [("a.txt", "Pump curve – rev B".encode("utf-8")), ("b.txt", "Pump curve – rev B".encode("utf-16")),
                       ("c.csv", "Pump curve – rev B".encode("cp1252")), ("d.md", b"\xef\xbb\xbfPump curve \xe2\x80\x93 rev B")]:
        (tmp_path / name).write_bytes(data)
        found = extract_text(tmp_path / name, CAP)
        assert (found.status, found.text) == ("ok", "Pump curve – rev B"), name


def test_pdf(tmp_path):
    make_pdf(tmp_path / "curve.pdf", "Design point 38 kW")
    found = extract_text(tmp_path / "curve.pdf", CAP)
    assert found.status == "ok" and "Design point 38 kW" in found.text


def test_docx_paragraphs_and_tables(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Installation procedure, rev 2")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Step 7.4", "Laser-align the coupling"
    document.save(tmp_path / "procedure.docx")
    found = extract_text(tmp_path / "procedure.docx", CAP)
    assert found.status == "ok"
    assert "Installation procedure, rev 2" in found.text and "Step 7.4 | Laser-align the coupling" in found.text


A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def rels(*items: tuple[str, str, str]) -> str:
    """A relationships part: (ID, type's last word, target), or an external link when the type is "hyperlink"."""
    out = "".join(f'<Relationship Id="{i}" Type="{RELS}/{kind}" Target="{target}"' + (' TargetMode="External"' if kind == "hyperlink" else "") + "/>"
                  for i, kind, target in items)
    return f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{out}</Relationships>'


def box(*paragraphs: str) -> str:
    """A text box: each paragraph's runs split at "|", and a line break where a paragraph has "/"."""
    def para(text: str) -> str:
        return "<a:p>" + "<a:br/>".join("".join(f"<a:r><a:t>{run}</a:t></a:r>" for run in line.split("|")) for line in text.split("/")) + "</a:p>"
    return f"<p:sp><p:nvSpPr/><p:txBody><a:bodyPr/>{''.join(para(p) for p in paragraphs)}</p:txBody></p:sp>"


def make_pptx(path: Path, slides: list[tuple[str, str | None]], order: list[int]) -> None:
    """A .pptx written by hand, with no package: each slide's shapes and its notes (or None). ``order`` is the deck's
    slide order, by slide number, which needn't be the order of the files."""
    parts = {
        "_rels/.rels": rels(("rId1", "officeDocument", "ppt/presentation.xml")),
        "ppt/presentation.xml": f"<p:presentation {P} {R}><p:sldIdLst>" + "".join(f'<p:sldId id="{255 + n}" r:id="rId{n}"/>' for n in order)
                                + "</p:sldIdLst></p:presentation>",
        "ppt/_rels/presentation.xml.rels": rels(*[(f"rId{n}", "slide", f"slides/slide{n}.xml") for n in range(1, len(slides) + 1)]),
    }
    for n, (shapes, notes) in enumerate(slides, 1):
        parts[f"ppt/slides/slide{n}.xml"] = f"<p:sld {A} {P}><p:cSld><p:spTree><p:nvGrpSpPr/>{shapes}</p:spTree></p:cSld></p:sld>"
        links = [("rId9", "hyperlink", "https://example.com/plan")]
        if notes is not None:
            links.append(("rId2", "notesSlide", f"../notesSlides/notesSlide{n}.xml"))
            image = '<p:sp><p:nvSpPr><p:nvPr><p:ph type="sldImg"/></p:nvPr></p:nvSpPr></p:sp>'
            body = f'<p:sp><p:nvSpPr><p:cNvPr/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{notes}</a:t></a:r></a:p></p:txBody></p:sp>'
            parts[f"ppt/notesSlides/notesSlide{n}.xml"] = f"<p:notes {A} {P}><p:cSld><p:spTree>{image}{body}</p:spTree></p:cSld></p:notes>"
        parts[f"ppt/slides/_rels/slide{n}.xml.rels"] = rels(*links)
    with zipfile.ZipFile(path, "w") as z:
        for name, text in parts.items():
            z.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + text)


def test_pptx_slides_tables_groups_and_notes(tmp_path):
    """Read with the standard library (python-pptx and its Pillow are gone): the deck's slide order, each shape's text,
    tables row by row, text inside groups, line breaks, and the speaker notes."""
    table = ("<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
             "<a:tr><a:tc><a:txBody><a:p><a:r><a:t>Task</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Owner</a:t></a:r></a:p></a:txBody></a:tc></a:tr>"
             "<a:tr><a:tc><a:txBody><a:p><a:r><a:t>Rigging plan</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Mike</a:t></a:r></a:p></a:txBody></a:tc></a:tr>"
             "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>")
    slides = [
        (box("Turnaround plan") + box("Isolate the reactor", "Laser-|align/the coupling"), "Crane booked for 20 Oct"),
        (box("Who does what") + table + f"<p:grpSp><p:nvGrpSpPr/><p:grpSpPr/>{box('Inside a group')}</p:grpSp>", None),
    ]
    make_pptx(tmp_path / "plan.pptx", slides, order=[2, 1])
    found = extract_text(tmp_path / "plan.pptx", CAP)
    assert found.status == "ok"
    assert found.text.split("\n") == ["Slide 1", "Who does what", "Task | Owner", "Rigging plan | Mike", "Inside a group",
                                      "Slide 2", "Turnaround plan", "Isolate the reactor", "Laser-align", "the coupling",
                                      "Notes: Crane booked for 20 Oct"]


def test_a_pptx_built_to_trick_the_reader_is_refused(tmp_path, monkeypatch):
    from docrepo import extract

    make_pptx(tmp_path / "plan.pptx", [(box("Turnaround plan"), None)], order=[1])
    with zipfile.ZipFile(tmp_path / "plan.pptx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["ppt/slides/slide1.xml"] = b'<!DOCTYPE x [<!ENTITY e "boom">]>' + parts["ppt/slides/slide1.xml"].split(b"?>", 1)[1]
    with zipfile.ZipFile(tmp_path / "dtd.pptx", "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    refused = extract_text(tmp_path / "dtd.pptx", CAP)
    assert refused.status == "error" and "ValueError" in refused.note
    monkeypatch.setattr(extract, "MAX_PART_BYTES", 100)
    assert extract_text(tmp_path / "plan.pptx", CAP).status == "error"  # a part bigger than the limit


def test_xlsx_sheets_and_cells(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Area 3"
    sheet.append(["Item", "Finding", "Status"])
    sheet.append(["Walkway rails", "Paint faded", "Open"])
    book.save(tmp_path / "checklist.xlsx")
    found = extract_text(tmp_path / "checklist.xlsx", CAP)
    assert found.text.split("\n") == ["Sheet: Area 3", "Item | Finding | Status", "Walkway rails | Paint faded | Open"]


def test_eml(tmp_path):
    (tmp_path / "quote.eml").write_bytes(b"From: Calder <q@example.com>\r\nSubject: Quote 4471\r\n\r\nTwo days on site.\r\n")
    found = extract_text(tmp_path / "quote.eml", CAP)
    assert "Subject: Quote 4471" in found.text and "Two days on site." in found.text


def test_long_text_is_cut_and_marked_partial(tmp_path):
    (tmp_path / "long.txt").write_text("x" * 50, encoding="utf-8")
    found = extract_text(tmp_path / "long.txt", 20)
    assert (found.status, len(found.text)) == ("partial", 20)


def test_files_it_cant_read(tmp_path):
    (tmp_path / "photo.jpg").write_bytes(b"\xff\xd8\xff")
    assert extract_text(tmp_path / "photo.jpg", CAP).status == "none"
    (tmp_path / "broken.docx").write_bytes(b"not a zip")
    broken = extract_text(tmp_path / "broken.docx", CAP)
    assert broken.status == "error" and "couldn't read it" in broken.note
    (tmp_path / "empty.txt").write_text("  ", encoding="utf-8")
    assert extract_text(tmp_path / "empty.txt", CAP).status == "none"
    assert extract_text(tmp_path / "missing.txt", CAP).status == "error"


def test_file_info(tmp_path):
    (tmp_path / "a.txt").write_bytes(b"abc")
    info = file_info(tmp_path / "a.txt")
    assert info.size == 3 and info.sha256 == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"


def test_local_paths_stay_as_they_are(tmp_path):
    assert to_unc(str(tmp_path / "a.txt")) == str(tmp_path / "a.txt")
    assert to_unc("\\\\fileserver\\eng\\a.pdf") == "\\\\fileserver\\eng\\a.pdf"
